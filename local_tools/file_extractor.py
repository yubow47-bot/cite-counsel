"""Full document extraction with stable offsets and honest coverage.

Model-facing previews may be short; raw_text is never silently shortened.
PDF OCR is bounded per call and can be continued using one-based page ranges.
"""

import os
import tempfile

SCANNED_THRESHOLD = 50
MAX_OCR_PAGES = 3


def extract_from_file(file_path: str, *, start_page: int | None = None,
                      end_page: int | None = None) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    if start_page is not None or end_page is not None:
        if ext != ".pdf":
            raise ValueError("Page ranges are supported only for PDF files.")
        return _extract_pdf(file_path, start_page=start_page, end_page=end_page)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".pptx":
        return _extract_pptx(file_path)
    if ext == ".xlsx":
        return _extract_xlsx(file_path)
    if ext in (".jpg", ".jpeg", ".png", ".webp"):
        from llm_api.vision import extract_from_image
        result = extract_from_image(file_path)
        if result.get("raw_text"):
            text, segments = assemble_segments([{"unit": 1, "kind": "image", "text": result["raw_text"]}])
            result.update(raw_text=text, segments=segments)
        return result
    return {"raw_input": f"Unsupported file type: {ext}", "error": f"Unsupported file type: {ext}"}


def head_tail(text: str, head: int, tail: int) -> str:
    """A display preview only; never use this as the stored document text."""
    if len(text) <= head + tail:
        return text
    return text[:head].rstrip() + "\n[…]\n" + text[-tail:].lstrip()


def assemble_segments(parts: list[dict]) -> tuple[str, list[dict]]:
    """Join located text while retaining half-open offsets into raw_text."""
    texts, segments, cursor = [], [], 0
    for part in parts:
        text = str(part.get("text") or "").strip()
        if not text:
            continue
        if texts:
            cursor += 1
        segment = {key: value for key, value in part.items() if key != "text"}
        segment.update(start=cursor, end=cursor + len(text))
        texts.append(text)
        segments.append(segment)
        cursor += len(text)
    return "\n".join(texts), segments


def _coverage(format: str, total: int, processed: list[int], missing: list[int] | None = None,
              **details) -> dict:
    missing = missing if missing is not None else sorted(set(range(1, total + 1)) - set(processed))
    return {"status": "partial" if missing else "complete", "format": format,
            "total_units": total, "processed_units": sorted(set(processed)),
            "missing_units": sorted(set(missing)), **details}


def _extract_docx(file_path: str) -> dict:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc, parts = Document(file_path), []
    props = doc.core_properties
    paragraph, table = 0, 0
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph += 1
            parts.append({"unit": paragraph, "kind": "paragraph", "paragraph": paragraph, "text": block.text})
        elif isinstance(block, Table):
            table += 1
            for row_number, row in enumerate(block.rows, 1):
                for column, cell in enumerate(row.cells, 1):
                    parts.append({"kind": "table_cell", "table": table, "row": row_number,
                                  "column": column, "text": _cell_text(cell)})
    text, segments = assemble_segments(parts)
    return {"title": props.title or _guess_title(text), "author": props.author or "",
            "date": str(props.created.date()) if props.created else "",
            "publisher": props.last_modified_by or "", "raw_text": text, "segments": segments,
            "coverage": _coverage("docx", paragraph, list(range(1, paragraph + 1)),
                                  tables=table, scope="main_body_and_tables",
                                  excluded_parts=["headers", "footers", "footnotes", "comments", "embedded_objects"])}


def _cell_text(cell) -> str:
    """Include nested tables, which cell.text alone does not read."""
    from docx.table import Table
    return "\n".join("\n".join(_cell_text(c) for row in block.rows for c in row.cells)
                     if isinstance(block, Table) else block.text for block in cell.iter_inner_content())


def _page_range(total: int, start_page: int | None, end_page: int | None) -> list[int]:
    first, last = start_page or 1, end_page or total
    if total == 0:
        return []
    if first < 1 or last < first or last > total:
        raise ValueError(f"PDF page range must be within 1–{total}.")
    return list(range(first, last + 1))


def _has_unread_raster(page) -> bool:
    """Detect a scanned body hidden behind a readable header or thin OCR layer.

    Small logos do not trigger paid OCR. A large raster region needs a text
    layer spread through that region, rather than a document-wide character
    threshold. This is a routing heuristic, not proof of transcription quality.
    """
    images = getattr(page, "images", [])
    if not isinstance(images, list) or not images:
        return False
    width, height = float(page.width), float(page.height)
    if width <= 0 or height <= 0:
        return True
    chars = getattr(page, "chars", [])
    for image in images:
        left, right = float(image["x0"]), float(image["x1"])
        top, bottom = float(image["top"]), float(image["bottom"])
        if (right - left) * (bottom - top) < width * height * 0.25:
            continue
        inside = [char for char in chars if char.get("text", "").strip()
                  and left <= (char["x0"] + char["x1"]) / 2 <= right
                  and top <= (char["top"] + char["bottom"]) / 2 <= bottom]
        if len(inside) < 200 or max(char["bottom"] for char in inside) - min(char["top"] for char in inside) < (bottom - top) / 2:
            return True
    return False


def _extract_pdf(file_path: str, *, start_page: int | None = None,
                 end_page: int | None = None) -> dict:
    import pdfplumber

    parts, errors = [], []
    with pdfplumber.open(file_path) as pdf:
        total = len(pdf.pages)
        selected = _page_range(total, start_page, end_page)
        for number in selected:
            try:
                page = pdf.pages[number - 1]
                text = page.extract_text() or ""
                parts.append({"unit": number, "kind": "page", "page": number,
                              "method": "text_layer", "text": text,
                              "complete": bool(text.strip()) and not _has_unread_raster(page)})
            except Exception:
                errors.append({"page": number, "reason": "text_extraction_failed"})
    if not errors and sum(len(p["text"].strip()) for p in parts) < SCANNED_THRESHOLD:
        if start_page is None and end_page is None:
            scanned = _extract_pdf_scanned(file_path)
        else:
            scanned = _extract_pdf_scanned(file_path, start_page=start_page, end_page=end_page)
        # A failed OCR attempt must not discard the short text layer that was
        # already readable (e.g. a cover page or a few labels on a scan).
        if isinstance(scanned, dict):
            existing = {s["page"]: {**s, "text": scanned["raw_text"][s["start"]:s["end"]]}
                        for s in scanned.get("segments", [])}
            for part in parts:
                if part["text"].strip() and part["page"] not in existing:
                    existing[part["page"]] = {**part, "complete": False}
            scanned["raw_text"], scanned["segments"] = assemble_segments([existing[p] for p in sorted(existing)])
            scanned["title"] = scanned.get("title") or _guess_title(scanned["raw_text"])
        return scanned
    by_page = {p["page"]: p for p in parts if p["text"].strip()}
    needs_ocr = [p["page"] for p in parts if not p["complete"]]
    if needs_ocr:
        ocr = _ocr_pages(file_path, needs_ocr[:MAX_OCR_PAGES])
        by_page.update({p["page"]: p for p in ocr["parts"]})
        errors.extend(ocr["errors"])
    readable = [by_page[number] for number in sorted(by_page)]
    text, segments = assemble_segments(sorted(readable, key=lambda p: p["page"]))
    processed = [p["page"] for p in readable if p.get("complete", True)]
    coverage = _coverage("pdf", total, processed, requested_pages=selected, errors=errors,
                         ocr_page_limit=MAX_OCR_PAGES, transcription_verified=False,
                         scope="text_layer_and_detected_scanned_pages",
                         raster_detection="large_image_text_layer_heuristic")
    return {"title": _guess_title(text), "raw_text": text, "segments": segments, "coverage": coverage}


def _ocr_pages(file_path: str, page_numbers: list[int]) -> dict:
    import fitz
    from llm_api.vision import extract_from_image

    parts, errors, titles = [], [], []
    with tempfile.TemporaryDirectory(prefix="legal-ocr-") as directory:
        with fitz.open(file_path) as doc:
            for number in page_numbers:
                try:
                    page = doc.load_page(number - 1)
                    w, h, scale = page.rect.width, page.rect.height, 200 / 72
                    if w * h * scale * scale > 16_000_000:
                        scale = (16_000_000 / (w * h)) ** 0.5
                    image_path = os.path.join(directory, f"page-{number}.png")
                    page.get_pixmap(matrix=fitz.Matrix(scale, scale)).save(image_path)
                    result = extract_from_image(image_path)
                    if result.get("error") or not result.get("raw_text"):
                        errors.append({"page": number, "reason": "vision_extraction_failed"})
                        continue
                    complete = result.get("coverage", {}).get("status") == "complete"
                    parts.append({"unit": number, "kind": "page", "page": number, "method": "vision",
                                  "complete": complete, "text": result["raw_text"]})
                    if not complete:
                        errors.append({"page": number, "reason": "vision_response_incomplete"})
                    titles.append(result.get("title") or result.get("page_title") or "")
                except Exception:
                    errors.append({"page": number, "reason": "page_render_or_extraction_failed"})
    return {"parts": parts, "errors": errors, "titles": titles}


def _extract_pdf_scanned(file_path: str, *, start_page: int | None = None,
                         end_page: int | None = None) -> dict:
    import fitz
    with fitz.open(file_path) as doc:
        total = len(doc)
    selected = _page_range(total, start_page, end_page)
    result = _ocr_pages(file_path, selected[:MAX_OCR_PAGES])
    text, segments = assemble_segments(result["parts"])
    processed = [p["page"] for p in result["parts"] if p["complete"]]
    return {"title": next((t for t in result["titles"] if t), _guess_title(text)),
            "raw_text": text, "segments": segments,
            "coverage": _coverage("pdf", total, processed, requested_pages=selected,
                                  errors=result["errors"], ocr_page_limit=MAX_OCR_PAGES,
                                  method="vision", transcription_verified=False)}


def _extract_pptx(file_path: str) -> dict:
    from pptx import Presentation
    prs, parts = Presentation(file_path), []
    props = prs.core_properties
    for number, slide in enumerate(prs.slides, 1):
        for shape_number, shape in enumerate(slide.shapes, 1):
            parts.extend(_shape_parts(shape, number, str(shape_number)))
    text, segments = assemble_segments(parts)
    return {"title": props.title or _guess_title(text), "author": props.author or "",
            "date": str(props.created.date()) if props.created else "", "raw_text": text,
            "segments": segments, "coverage": _coverage("pptx", len(prs.slides),
                                                        list(range(1, len(prs.slides) + 1)),
                                                        scope="slide_text_and_tables",
                                                        excluded_parts=["images", "speaker_notes", "chart_data"])}


def _shape_parts(shape, slide: int, shape_path: str) -> list[dict]:
    parts = []
    location = {"unit": slide, "slide": slide, "shape": shape_path}
    if shape.has_text_frame and shape.text.strip():
        parts.append({**location, "kind": "slide_text", "text": shape.text})
    if shape.has_table:
        for number, row in enumerate(shape.table.rows, 1):
            parts.append({**location, "kind": "table_row", "row": number,
                          "text": "\t".join(cell.text for cell in row.cells)})
    if hasattr(shape, "shapes"):
        for number, child in enumerate(shape.shapes, 1):
            parts.extend(_shape_parts(child, slide, f"{shape_path}.{number}"))
    return parts


def _extract_xlsx(file_path: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    try:
        props, parts = wb.properties, []
        for number, sheet in enumerate(wb.worksheets, 1):
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if not any(c is not None for c in row):
                    continue
                parts.append({"unit": number, "kind": "sheet_row", "sheet": sheet.title,
                              "row": row_number, "text": "\t".join("" if c is None else str(c) for c in row)})
        text, segments = assemble_segments(parts)
        return {"title": props.title or _guess_title(text), "author": props.creator or "",
                "date": str(props.created.date()) if props.created else "", "raw_text": text,
                "segments": segments, "coverage": _coverage("xlsx", len(wb.worksheets),
                                                            list(range(1, len(wb.worksheets) + 1)),
                                                            scope="cell_values_and_formulas", formulas="preserved_not_evaluated")}
    finally:
        wb.close()


def _guess_title(text: str) -> str:
    for line in text.splitlines():
        if len(line.strip()) > 5:
            return line.strip()[:100]
    return ""

