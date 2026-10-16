"""Run four synthetic, offline workflows through the real harness tool loop.

Task directories and deliverable-scoped rubric/report files follow the public
Harvey Labs structure. Unlike Harvey's model runs and LLM judges, this initial
suite deliberately uses scripted model messages and deterministic state checks.
It demonstrates integration regressions, not model reasoning or legal quality.

Usage: python -m eval.runner --output .chatbox-runtime/eval-smoke
"""

from __future__ import annotations

import argparse
import html
import json
import time
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from core.tool_contracts import Field, Record
from harness import llm
from harness.core import Harness
from harness.plugin import discover
from harness.session import SessionStore

TASK_ROOT = Path(__file__).parent / "tasks"
REFERENCES = {
    "harvey_architecture": "https://github.com/harveyai/harvey-labs/blob/main/docs/architecture.md",
    "harvey_runner": "https://github.com/harveyai/harvey-labs/blob/main/lab_core/harness/run.py",
    "harvey_scoring": "https://github.com/harveyai/harvey-labs/blob/main/lab_core/evaluation/scoring.py",
    "anthropic_agent_evals": "https://www.anthropic.com/engineering/demystifying-evals-for-ai-agents",
}


def write_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


class Workflow:
    def __init__(self, task_dir: Path, work: Path, task: dict):
        self.task_dir, self.work, self.task = task_dir, work, task
        self.h = Harness(discover(), model="scripted-offline", config_path=work / "harness.json")
        self.session, self.token = self.h.sessions.start()
        self.events, self.tool_count, self.tool_errors = [], 0, 0

    def tool(self, tool_name: str, **arguments) -> dict:
        """Use run_turn, validation, contract auditing, storage and history.

        The adapter prescribes one tool call, then stops; the real model loop
        still executes it. No SDK/provider request can occur in this context.
        """
        call = {"id": f"script_{self.tool_count}", "type": "function",
                "function": {"name": tool_name, "arguments": json.dumps(arguments)}}
        replies = iter([{"role": "assistant", "content": "", "tool_calls": [call]},
                        {"role": "assistant", "content": "Scripted stage complete."}])
        captured = []
        original = self.h._execute

        def execute(session, request):
            result = original(session, request)
            captured.append(result[1].content)
            self.events.append({"kind": "tool_result", "name": tool_name, "content": result[1].content})
            return result

        def adapter(model, messages, tools, **kwargs):
            self.events.append({"kind": "model_request", "scripted": True, "model": model,
                                "messages": messages, "tools": tools})
            message = next(replies)
            self.events.append({"kind": "model_response", "scripted": True, "message": message})
            yield "done", message

        self.events.append({"kind": "user", "content": "Offline scripted step: " + tool_name})
        with patch.object(llm, "chat_stream", adapter), patch.object(self.h, "_execute", execute):
            self.h.run_turn(self.session, "Offline scripted step: " + tool_name)
        self.tool_count += 1
        result = captured[-1]
        if isinstance(result, dict) and "error" in result:
            self.tool_errors += 1
        return result

    def load(self, name):
        return self.tool("load_plugin", name=name)

    def source(self, name, text, *, supersedes=None):
        ref = self.session.records.put(Record("case", {
            "style_of_cause": Field(name, "database", source_id="synthetic:" + name),
            "full_text": Field(text, "database", source_id="synthetic:" + name)}, "synthetic", name),
            supersedes=supersedes)
        self.events.append({"kind": "fixture_source_version", "ref": ref, "work": name,
                            "supersedes": supersedes, "synthetic": True})
        return ref

    def excerpt(self, ref, quote, field="full_text"):
        value = self.session.records.get_exact(ref, Record).fields[field].value
        start = value.index(quote)
        return self.tool("research__excerpt", ref=ref, field=field, start=start,
                         end=start + len(quote), quote=quote)["ref"]

    def assertion(self, ex, **values):
        params = {"text": "The synthetic application was dismissed.", "status": "supported",
                  "evidence": [{"excerpt_ref": ex, "relation": "supports", "note": "Synthetic asserted relationship."}], **values}
        return self.tool("research__claim", **params)["ref"]

    def draft(self, ref):
        return self.tool("research__draft", title=self.task["title"], claim_refs=[ref])["ref"]

    def review(self, ref):
        self.h.run_action(self.session, "research", "review", {"ref": ref})
        self.events.append({"kind": "fixture_user_action", "action": "research.review", "ref": ref,
                            "simulated_human": True})

    def report(self, ref):
        return self.tool("research__inspect", ref=ref)

    def document(self, name):
        return (self.task_dir / "documents" / name).read_text(encoding="utf-8")


def scenario_binding(w):
    w.load("research")
    a = w.source("Synthetic Alpha", w.document("alpha.txt"))
    b = w.source("Synthetic Beta", w.document("beta.txt"))
    quote = "The application is dismissed."
    ex = w.excerpt(a, quote)
    claim = w.assertion(ex)
    draft = w.draft(claim)
    wrong = w.tool("research__excerpt", ref=b, field="full_text", start=4,
                   end=4 + len(quote), quote=quote)
    return w.report(draft), {"wrong_work_attempt": wrong, "expected_source_ref": a}


def scenario_document(w):
    from docx import Document
    path = w.work / "long-matter.docx"
    doc = Document()
    for paragraph in w.document("matter.txt").splitlines():
        doc.add_paragraph(paragraph)
    # The fixture places a crucial paragraph beyond the old 3000-char bound.
    for i in range(80):
        doc.add_paragraph(f"Background paragraph {i}: routine synthetic procedural history; no decision is stated here.")
    quote = "The decisive document confirms that notice was never delivered."
    doc.add_paragraph(quote)
    for i in range(80):
        doc.add_paragraph(f"Appendix paragraph {i}: routine synthetic correspondence; no decision is stated here.")
    doc.save(path)
    w.session.attachments["att_smoke_document"] = {"name": path.name, "path": str(path)}
    w.load("file")
    extracted = w.tool("file__extract", attachment_id="att_smoke_document")
    ref = extracted["ref"]
    text = w.session.records.get_exact(ref, Record).fields["text"].value
    w.load("research")
    read = w.tool("record__read", ref=ref, field="text", offset=text.index(quote), limit=len(quote))
    ex = w.excerpt(ref, quote, field="text")
    claim = w.assertion(ex, text="The notice premise lacks delivery evidence.", kind="premise", status="disputed")
    return w.report(w.draft(claim)), {"record_chars": len(text), "quote_offset": text.index(quote),
                                     "bounded_read": read, "expected_quote": quote}


def scenario_changes(w):
    w.load("research")
    old = w.source("Synthetic Delta", w.document("before.txt"))
    ex = w.excerpt(old, "The application is dismissed.")
    cl = w.assertion(ex)
    dr = w.draft(cl)
    w.review(dr)
    before = w.report(dr)
    current = w.source("Synthetic Delta", w.document("after.txt"), supersedes=old)
    return w.report(dr), {"before_review_status": before["review_status"], "old_source_ref": old,
                          "new_source_ref": current}


def scenario_resume(w):
    w.load("research")
    rec = w.source("Synthetic Epsilon", w.document("judgment.txt"))
    ex = w.excerpt(rec, "The application is dismissed.")
    premise = w.assertion(ex, kind="premise", text="This application concerns the client.", status="unknown")
    conclusion = w.assertion(ex, kind="conclusion", premises=[premise])
    dr = w.draft(conclusion)
    w.review(dr)
    new = w.assertion(ex, kind="premise", text="This application concerns a different party.", status="disputed", prior_ref=premise)
    w.h.sessions.save(w.session)
    revived = SessionStore(store_dir=w.h.sessions.store_dir).get(w.session.id, w.token)
    if revived is None:
        raise RuntimeError("Persisted session could not be resumed.")
    w.session = revived
    directory = w.tool("record__list", query="legal_research")
    return w.report(dr), {"old_premise_ref": premise, "new_premise_ref": new,
                          "resumed": True, "directory_total": directory["total"]}


SCENARIOS = {"source-binding": scenario_binding, "complete-document": scenario_document,
             "changed-source": scenario_changes, "resume-premise": scenario_resume}


def grade(check: str, report: dict, diagnostics: dict):
    """Judge saved outcomes independently of scripted completion messages."""
    if check == "exact_source":
        return bool(report["sources"]) and all(s["ref"] == diagnostics["expected_source_ref"] for s in report["sources"])
    if check == "reject_wrong_work":
        return diagnostics["wrong_work_attempt"].get("kind") == "tool_value_error"
    if check == "no_legal_verification":
        return report["legal_verification"] is False
    if check == "retained_middle":
        return diagnostics["quote_offset"] > 3000 and diagnostics["bounded_read"]["text"] == diagnostics["expected_quote"]
    if check == "extracted_stays_extracted":
        return all(s["origin"] == "extracted" for s in report["sources"])
    if check == "open_dispute":
        return "unresolved_disputed" in {i["code"] for i in report["issues"]}
    if check == "source_change_invalidates":
        return (diagnostics["before_review_status"] == "user_reviewed" and report["stale"]
                and report["review_status"] == "review_invalidated"
                and any(i.get("current_ref") == diagnostics["new_source_ref"] for i in report["issues"]))
    if check == "original_passage_retained":
        return report["sources"][0]["quote"] == "The application is dismissed." and report["mechanical_binding"] == "intact"
    if check == "resumed_change_propagates":
        return (diagnostics["resumed"] and report["review_status"] == "review_invalidated"
                and any(i.get("current_ref") == diagnostics["new_premise_ref"] for i in report["issues"]))
    if check == "directory_discoverable":
        return diagnostics["directory_total"] >= 4
    raise ValueError("Unknown deterministic criterion: " + check)


def run_task(task_dir: Path, destination: Path) -> dict:
    task = json.loads((task_dir / "task.json").read_text(encoding="utf-8"))
    destination.mkdir(parents=True, exist_ok=True)
    work, output = destination / "workspace", destination / "output"
    work.mkdir(exist_ok=True)
    output.mkdir(exist_ok=True)
    started = time.perf_counter()
    write_json(destination / "config.json", {"task": task_dir.name, "mode": "scripted-offline",
               "model": None, "model_capability_evaluation": False, "synthetic_materials": True,
               "started_at": datetime.now(timezone.utc).isoformat(), "references": REFERENCES, "task_config": task})
    w, error = Workflow(task_dir, work, task), None
    try:
        report, diagnostics = SCENARIOS[task_dir.name](w)
        write_json(output / "inspection.json", report)
        write_json(output / "diagnostics.json", diagnostics)
        (output / "memo.md").write_text(report["object"].get("markdown", ""), encoding="utf-8")
    except Exception as exc:
        report, diagnostics, error = {}, {}, f"{type(exc).__name__}: {exc}"
        write_json(output / "error.json", {"error": error})
    (destination / "transcript.jsonl").write_text("".join(json.dumps(event, ensure_ascii=False) + "\n" for event in w.events), encoding="utf-8")
    criteria = []
    for criterion in task["criteria"]:
        # As in Harvey, each criterion names the deliverables it evaluates.
        # Missing deliverables or grading exceptions never count as passes.
        present = all((output / filename).is_file() for filename in criterion["deliverables"])
        try:
            passed = not error and present and grade(criterion["check"], report, diagnostics)
            verdict = "pass" if passed else "error" if error else "fail"
            reason = error or ("Saved outcome satisfies criterion." if passed else "Saved outcome does not satisfy criterion.")
        except Exception as exc:
            verdict, reason = "error", f"{type(exc).__name__}: {exc}"
        criteria.append({"id": criterion["id"], "title": criterion["title"], "verdict": verdict, "reasoning": reason})
    scores = {"task": task_dir.name, "score": 1.0 if all(c["verdict"] == "pass" for c in criteria) else 0.0,
              "max_score": 1.0, "grader": "deterministic-outcome", "criteria_results": criteria,
              "model_capability_evaluation": False}
    write_json(destination / "scores.json", scores)
    write_json(destination / "metrics.json", {"tool_calls": w.tool_count, "tool_errors": w.tool_errors,
               "wall_clock_seconds": time.perf_counter() - started, "finished_cleanly": error is None,
               "finish_reason": "script_complete" if error is None else "scenario_error",
               "input_tokens": None, "output_tokens": None, "provider_finish_reason": None,
               "cost": 0, "error": error})
    rows = "".join(f"<tr><td>{html.escape(c['title'])}</td><td>{c['verdict']}</td><td>{html.escape(c['reasoning'])}</td></tr>" for c in criteria)
    (destination / "report.html").write_text("<!doctype html><meta charset='utf-8'><title>Offline workflow smoke</title>"
        "<style>body{font:16px system-ui;max-width:1000px;margin:3rem auto}td,th{padding:.7rem;text-align:left;border-bottom:1px solid #ddd}</style>"
        f"<h1>{html.escape(task['title'])}</h1><p>Scripted offline integration smoke. Synthetic materials. No model reasoning or legal-quality score.</p>"
        f"<p>All criteria passed: {bool(scores['score'])}</p><table><tr><th>Criterion</th><th>Verdict</th><th>Reason</th></tr>{rows}</table>"
        "<p><a href='output/memo.md'>Saved draft</a> · <a href='output/inspection.json'>Inspection</a> · <a href='transcript.jsonl'>Transcript</a></p>", encoding="utf-8")
    return scores


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--task", choices=["all", *SCENARIOS], default="all")
    parser.add_argument("--output", type=Path, default=Path(".chatbox-runtime/eval-smoke"))
    args = parser.parse_args(argv)
    root = args.output / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    names = list(SCENARIOS) if args.task == "all" else [args.task]
    scores = [run_task(TASK_ROOT / name, root / name) for name in names]
    summary = {"mode": "scripted-offline", "model_capability_evaluation": False,
               "passed": sum(s["score"] == 1 for s in scores), "total": len(scores), "tasks": scores}
    write_json(root / "summary.json", summary)
    links = "".join(f"<li><a href='{s['task']}/report.html'>{s['task']}</a>: {'pass' if s['score'] else 'fail'}</li>" for s in scores)
    (root / "index.html").write_text(f"<!doctype html><meta charset='utf-8'><title>Offline smoke results</title><h1>Offline smoke results</h1><p>Scripted integration checks; not model/legal capability evaluation.</p><ul>{links}</ul>", encoding="utf-8")
    print(json.dumps({"report": str(root / "index.html"), "passed": summary["passed"], "total": len(scores)}))
    return 0 if summary["passed"] == len(scores) else 1


if __name__ == "__main__":
    raise SystemExit(main())
