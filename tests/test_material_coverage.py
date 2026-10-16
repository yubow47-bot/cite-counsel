"""Material regressions: late evidence survives extraction, storage and continuation."""
import json
from unittest.mock import patch

from local_tools.file_extractor import extract_from_file
from llm_api.vision import _align_fields


def _located(result, needle):
    segment = next(s for s in result["segments"] if needle in result["raw_text"][s["start"]:s["end"]])
    assert 0 <= segment["start"] < segment["end"] <= len(result["raw_text"])
    return segment


def test_long_docx_retains_middle_evidence_and_tables_in_document_order(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("Opening facts " * 500)
    document.add_paragraph("CRITICAL-MIDDLE-EVIDENCE")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TABLE-EXCEPTION"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED-EVIDENCE"
    document.add_paragraph("Conclusion " * 500)
    path = tmp_path / "long.docx"
    document.save(path)
    result = extract_from_file(str(path))
    assert len(result["raw_text"]) > 9000
    assert _located(result, "CRITICAL-MIDDLE-EVIDENCE")["paragraph"] == 2
    assert _located(result, "TABLE-EXCEPTION")["table"] == 1
    assert "NESTED-EVIDENCE" in result["raw_text"]
    assert result["raw_text"].index("TABLE-EXCEPTION") < result["raw_text"].index("Conclusion")
    assert result["coverage"]["scope"] == "main_body_and_tables"


def test_pdf_retains_evidence_after_page_five_and_reports_offsets(tmp_path):
    import fitz
    path = tmp_path / "long.pdf"
    with fitz.open() as document:
        for number in range(1, 8):
            page = document.new_page()
            page.insert_text((72, 72), f"Page {number}: " + "ordinary facts " * 7)
            if number == 7:
                page.insert_text((72, 110), "FINAL-PAGE-EXCEPTION")
        document.save(path)
    result = extract_from_file(str(path))
    assert _located(result, "FINAL-PAGE-EXCEPTION")["page"] == 7
    assert result["coverage"]["processed_units"] == list(range(1, 8))
    assert result["coverage"]["missing_units"] == []


def test_pdf_failed_ocr_retains_already_readable_short_text_layer(tmp_path):
    import fitz
    path = tmp_path / "short.pdf"
    with fitz.open() as document:
        document.new_page().insert_text((72, 72), "Visible label")
        document.save(path)
    with patch("llm_api.vision.extract_from_image", return_value={"error": "Budget unavailable"}):
        result = extract_from_file(str(path))
    assert "Visible label" in result["raw_text"]
    assert result["coverage"]["status"] == "partial"
    assert result["coverage"]["missing_units"] == [1]


def test_pptx_retains_late_slides_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    document = Presentation()
    for number in range(8):
        slide = document.slides.add_slide(document.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text = f"Slide {number + 1}: " + "context " * 100
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "LATE-SLIDE-EXCEPTION"
    path = tmp_path / "long.pptx"
    document.save(path)
    result = extract_from_file(str(path))
    assert len(result["raw_text"]) > 3000
    assert _located(result, "LATE-SLIDE-EXCEPTION")["slide"] == 8


def test_xlsx_retains_rows_after_fifty_and_formulas(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    for number in range(1, 121):
        sheet.append([number, "facts " * 15])
    sheet.cell(110, 2).value = "ROW-110-EXCEPTION"
    sheet.cell(111, 3).value = "=A110+1"
    path = tmp_path / "long.xlsx"
    workbook.save(path)
    result = extract_from_file(str(path))
    assert _located(result, "ROW-110-EXCEPTION")["row"] == 110
    assert "=A110+1" in result["raw_text"]
    assert result["coverage"]["formulas"] == "preserved_not_evaluated"


def test_vision_preserves_long_text_maps_title_and_exposes_incomplete_finish():
    text = "transcribed body " * 600
    complete = _align_fields({"page_title": "Judgment", "raw_text": text, "_vision_finish_reason": "stop"})
    assert complete["title"] == "Judgment"
    assert complete["raw_text"] == text.strip()
    assert complete["coverage"]["status"] == "complete"
    truncated = _align_fields({"raw_text": text, "_vision_finish_reason": "length"})
    assert truncated["coverage"]["status"] == "partial"
    assert truncated["coverage"]["missing_units"] == [1]


def test_scanned_pdf_can_be_continued_without_losing_earlier_pages(tmp_path):
    import fitz
    from harness.core import Harness
    from harness.plugin import discover
    from harness.session import Context
    from plugins import file

    path = tmp_path / "scan.pdf"
    with fitz.open() as document:
        for _ in range(6):
            document.new_page()
        document.save(path)
    harness = Harness(discover(), model="m")
    session, _ = harness.sessions.start()
    session.attachments["att_scan123"] = {"path": str(path), "name": "scan.pdf"}
    context = Context(session, "file", harness)

    def fake_vision(image_path):
        from pathlib import Path
        number = int(Path(image_path).stem.split("-")[1])
        return {"title": "Scan", "raw_text": f"Evidence on page {number}",
                "coverage": {"status": "complete"}}

    with patch("llm_api.vision.extract_from_image", side_effect=fake_vision) as vision:
        first = file.extract(context, file.ExtractParams(attachment_id="att_scan123"))
        assert first.content["coverage"]["missing_units"] == [4, 5, 6]
        assert vision.call_count == 3
        second = file.extract(context, file.ExtractParams(attachment_id="att_scan123", start_page=4, end_page=6))
        assert vision.call_count == 6
    record = session.records.get(second.content["ref"])
    assert all(f"Evidence on page {number}" in record.fields["text"].value for number in range(1, 7))
    assert second.content["coverage"]["status"] == "complete"
    assert second.content["replaces"] == first.content["ref"]
    assert [s["page"] for s in json.loads(record.fields["extraction_segments"].value)] == list(range(1, 7))


def test_plugin_stores_full_text_with_explicit_short_preview(tmp_path):
    from harness.core import Harness
    from harness.plugin import discover
    from harness.session import Context
    from plugins import file

    path = tmp_path / "material.docx"
    path.write_bytes(b"dummy attachment bytes")
    harness = Harness(discover(), model="m")
    session, _ = harness.sessions.start()
    session.attachments["att_long123"] = {"path": str(path), "name": "material.docx"}
    text = "ordinary evidence " * 700 + "FINAL-EXCEPTION"
    with patch("local_tools.file_extractor.extract_from_file", return_value={
        "title": "Material", "raw_text": text, "coverage": {"status": "complete"},
        "segments": [{"kind": "paragraph", "paragraph": 1, "start": 0, "end": len(text)}],
    }):
        result = file.extract(Context(session, "file", harness), file.ExtractParams(attachment_id="att_long123"))
    assert result.content["preview_truncated"] is True
    assert len(result.content["preview"]) == 3000
    assert session.records.get(result.content["ref"]).fields["text"].value == text
    assert result.content["text_chars"] == len(text)

